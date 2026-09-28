from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment

FILE = "Automated_Sales_Report.xlsx"

wb = load_workbook(FILE)

dashboard = wb["Dashboard"]
monthly = wb["Monthly Summary"]
product = wb["Product Summary"]

# Clear old dashboard chart area
for row in dashboard.iter_rows(min_row=1, max_row=40, min_col=4, max_col=20):
    for cell in row:
        cell.value = None

# Dashboard title
dashboard["A1"] = "Automated Sales Dashboard"
dashboard["A1"].font = Font(bold=True, size=16)

dashboard["A2"] = "KPI"
dashboard["B2"] = "Value"

for cell in dashboard[2]:
    cell.font = Font(bold=True)

# Monthly Sales Trend
line_chart = LineChart()
line_chart.title = "Monthly Sales Trend"
line_chart.y_axis.title = "Sales"
line_chart.x_axis.title = "Month"

data = Reference(
    monthly,
    min_col=3,
    min_row=1,
    max_row=monthly.max_row
)

categories = Reference(
    monthly,
    min_col=1,
    min_row=2,
    max_row=monthly.max_row
)

line_chart.add_data(data, titles_from_data=True)
line_chart.set_categories(categories)
line_chart.height = 9
line_chart.width = 18

dashboard.add_chart(line_chart, "D3")

# Product Sales
bar_chart = BarChart()
bar_chart.title = "Sales by Product"
bar_chart.y_axis.title = "Sales"
bar_chart.x_axis.title = "Product"

data = Reference(
    product,
    min_col=3,
    min_row=1,
    max_row=product.max_row
)

categories = Reference(
    product,
    min_col=1,
    min_row=2,
    max_row=product.max_row
)

bar_chart.add_data(data, titles_from_data=True)
bar_chart.set_categories(categories)
bar_chart.height = 9
bar_chart.width = 18

dashboard.add_chart(bar_chart, "D20")

# Format KPI values
dashboard["B4"].number_format = '?#,##0.00'
dashboard["B5"].number_format = '?#,##0.00'
dashboard["B6"].number_format = '?#,##0.00'
dashboard["B7"].number_format = '0.00%'

dashboard.column_dimensions["A"].width = 25
dashboard.column_dimensions["B"].width = 20

for row in range(3, 8):
    dashboard[f"A{row}"].font = Font(bold=True)

dashboard.freeze_panes = "A3"

wb.save(FILE)

print("Dashboard updated successfully!")
print("Charts added: Monthly Sales Trend + Sales by Product")
print(f"File: {FILE}")
